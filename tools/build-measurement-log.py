#!/usr/bin/env python3
"""Build final/measurement-log.xlsx: the record workbook for the final-report measurements.

Sheets: 說明, 血壓722, 臥站血壓, 夜間如廁, 洗澡濕度, 摘要.
Light-yellow cells are inputs; grey italic rows are examples and are not counted.
Every result is a formula, so the workbook updates as the measurements are typed in.

Usage:  python3 tools/build-measurement-log.py [--pdf]   (--pdf also writes final/measurement-sheets.pdf; needs LibreOffice Calc)
Then:   python3 <xlsx skill>/scripts/recalc.py final/measurement-log.xlsx   (fills cached values)
"""
import datetime
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "final", "measurement-log.xlsx")

FONT = "Arial"
F_BASE = Font(name=FONT, size=11)
F_TITLE = Font(name=FONT, size=16, bold=True)
F_NOTE = Font(name=FONT, size=10, color="595959")
F_HEAD = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=11, color="0000FF")
F_EXAMPLE = Font(name=FONT, size=11, italic=True, color="808080")
F_LINK = Font(name=FONT, size=11, color="008000")
F_BOLD = Font(name=FONT, size=11, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="11161D")
FILL_SUB = PatternFill("solid", fgColor="D9D9D9")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_EXAMPLE = PatternFill("solid", fgColor="F2F2F2")
FILL_ALERT = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
NOWRAP = Alignment(horizontal="left", vertical="center")

SRC_722 = "依據：衛生福利部國民健康署（2023-05-15）量血壓「722」原則，https://www.mohw.gov.tw/cp-16-74604-1.html"
SRC_OH = "步驟依據：CDC STEADI（2017）Measuring Orthostatic Blood Pressure，https://www.cdc.gov/steadi/media/pdfs/steadi-assessment-measuringbp-508.pdf；門檻依據：臺大醫院健康電子報（2023-10，第 215 期）姿勢性低血壓；Freeman R, et al. Clin Auton Res 2011;21(2):69–72"
SRC_NOC = "依據：Vaughan CP, et al. Int J Clin Pract 2010;64(5):577–583（每晚 ≥ 3 次夜尿與跌倒風險）；提醒條件為本設計的規則"


def style_range(ws, cells, font=None, fill=None, align=None, border=True, fmt=None):
    for row in ws[cells]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if align:
                c.alignment = align
            if border:
                c.border = BORDER
            if fmt:
                c.number_format = fmt


def title(ws, text, note, width_to):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = note
    ws["A2"].font = F_NOTE
    ws["A2"].alignment = LEFT
    ws.merge_cells(f"A2:{width_to}2")
    ws.row_dimensions[2].height = 46


def header(ws, row, labels, groups=None):
    """labels: list of column headers on `row`; groups: [(first_col, last_col, text)] on row-1."""
    if groups:
        for c1, c2, text in groups:
            ws.cell(row=row - 1, column=c1, value=text)
            if c2 > c1:
                ws.merge_cells(start_row=row - 1, start_column=c1, end_row=row - 1, end_column=c2)
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=row - 1, column=c)
                cell.font = F_HEAD
                cell.fill = FILL_HEAD
                cell.alignment = CENTER
                cell.border = BORDER
    for i, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=i, value=text)
        cell.font = F_BOLD if groups else F_HEAD
        cell.fill = FILL_SUB if groups else FILL_HEAD
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def summary_row(ws, row, label, formula, note=None, fmt=None):
    ws.cell(row=row, column=1, value=label).font = F_BOLD
    ws.cell(row=row, column=1).alignment = NOWRAP
    c = ws.cell(row=row, column=2, value=formula)
    c.font = F_BASE
    c.alignment = NOWRAP
    if fmt:
        c.number_format = fmt
    if note:
        n = ws.cell(row=row, column=3, value=note)
        n.font = F_NOTE
        n.alignment = NOWRAP


wb = Workbook()

# ---------------------------------------------------------------- 說明
ws = wb.active
ws.title = "說明"
ws["A1"] = "期末報告實測紀錄表（浴室安全監測系統）"
ws["A1"].font = F_TITLE
rows = [
    ("用途", "期末報告第 5、6 項（感測器訊號的收集、資料整理分析）的實際資料。量完把這個檔案交回，用來做期末報告的分析頁。"),
    ("怎麼填", "只填淺黃色格子（藍色字）。灰色斜體的「範例」列不會被計算，可以留著。計算欄、判斷欄與「摘要」工作表會自動更新。"),
    ("1. 血壓 722", "連續 7 天，早上起床後、晚上睡覺前各量一次，每次量兩遍（國健署 722 原則）。為了每天的數字可以互相比較，本紀錄固定坐著、量同一隻手臂，量完一遍隔一下再量第二遍。"),
    ("2. 臥站血壓", "每週 1 次（共 4 次）。平躺休息 5 分鐘後量一次，站起來後第 1 分鐘、第 3 分鐘各再量一次。站起來時旁邊要有人扶；如果頭暈，立刻坐下並停止測量。"),
    ("3. 夜間如廁", "連續 7 晚。記錄上床、起床時間，以及半夜起來上廁所的次數，起身時有沒有頭暈。可以請長輩早上回想，或由家人記錄。"),
    ("4. 洗澡濕度", "7 次洗澡。浴室放一個電子溫濕度計（大多數手機沒有濕度感測器）。記錄洗澡前的濕度、洗澡中看到的最高濕度、洗完 10 分鐘後的濕度，以及開始、結束時間。"),
    ("安全", "量測只是報告用的紀錄，不是診斷。血壓很高、站起來頭暈或身體不舒服，請停止並尋求醫療協助。開始前先徵得被量測者同意。"),
    ("個資", "交回檔案時，被量測者可以用「阿嬤」「同學」代稱，不需要寫真實姓名。"),
    ("格子顏色", "淺黃底、藍字＝要填的格子；灰底斜體＝範例；黑字＝公式；綠字＝從其他工作表帶入的結果；紅底＝超過文獻門檻。"),
]
for i, (k, v) in enumerate(rows, start=3):
    ws.cell(row=i, column=1, value=k).font = F_BOLD
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    c = ws.cell(row=i, column=2, value=v)
    c.font = F_BASE
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[i].height = 48
r = len(rows) + 4
ws.cell(row=r, column=1, value="文獻").font = F_BOLD
for j, src in enumerate([SRC_722, SRC_OH, SRC_NOC]):
    c = ws.cell(row=r + j, column=2, value=src)
    c.font = F_NOTE
    c.alignment = LEFT
    ws.row_dimensions[r + j].height = 44
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 100

# ---------------------------------------------------------------- 血壓722
ws = wb.create_sheet("血壓722")
title(ws, "血壓 7 天紀錄（722 原則）", "連續 7 天、早晚各一次、每次量 2 遍。單位：收縮壓、舒張壓 mmHg；脈搏 次/分。" + SRC_722, "N")
labels = ["天", "日期", "時段", "時間", "收縮壓", "舒張壓", "脈搏", "收縮壓", "舒張壓", "脈搏", "收縮壓", "舒張壓", "脈搏", "備註"]
header(ws, 4, labels, groups=[(1, 4, "量測"), (5, 7, "第 1 遍"), (8, 10, "第 2 遍"), (11, 13, "本次平均（自動）"), (14, 14, "")])
# example row
ex = ["範例", "2026-12-01", "早", "07:10", 128, 82, 72, 124, 80, 70, None, None, None, "範例列不計入"]
for i, v in enumerate(ex, start=1):
    ws.cell(row=5, column=i, value=v)
for col, src1, src2 in (("K", "E", "H"), ("L", "F", "I"), ("M", "G", "J")):
    ws[f"{col}5"] = f"=IF(COUNT({src1}5,{src2}5)=0,\"\",AVERAGE({src1}5,{src2}5))"
style_range(ws, "A5:N5", font=F_EXAMPLE, fill=FILL_EXAMPLE, align=CENTER)
FIRST, LAST = 6, 19
for k in range(14):
    r = FIRST + k
    ws.cell(row=r, column=1, value=f"第 {k // 2 + 1} 天")
    ws.cell(row=r, column=3, value="早" if k % 2 == 0 else "晚")
    for col, src1, src2 in (("K", "E", "H"), ("L", "F", "I"), ("M", "G", "J")):
        ws[f"{col}{r}"] = f"=IF(COUNT({src1}{r},{src2}{r})=0,\"\",AVERAGE({src1}{r},{src2}{r}))"
style_range(ws, f"A{FIRST}:A{LAST}", font=F_BASE, align=CENTER)
style_range(ws, f"B{FIRST}:J{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"N{FIRST}:N{LAST}", font=F_INPUT, fill=FILL_INPUT, align=LEFT)
style_range(ws, f"K{FIRST}:M{LAST}", font=F_BASE, align=CENTER, fmt="0.0")
style_range(ws, f"K5:M5", fmt="0.0")
ws.conditional_formatting.add(f"K{FIRST}:K{LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["130"], fill=FILL_ALERT))
ws.conditional_formatting.add(f"L{FIRST}:L{LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["80"], fill=FILL_ALERT))
dv = DataValidation(type="list", formula1='"早,晚"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"C{FIRST}:C{LAST}")
S = LAST + 2
rng_s, rng_d, rng_p = f"K{FIRST}:K{LAST}", f"L{FIRST}:L{LAST}", f"M{FIRST}:M{LAST}"
summary_row(ws, S, "完成次數", f"=COUNT({rng_s})", "共 14 次（7 天 × 早晚）")
summary_row(ws, S + 1, "平均收縮壓", f"=IF(COUNT({rng_s})=0,\"\",ROUND(AVERAGE({rng_s}),1))", "mmHg", "0.0")
summary_row(ws, S + 2, "平均舒張壓", f"=IF(COUNT({rng_d})=0,\"\",ROUND(AVERAGE({rng_d}),1))", "mmHg", "0.0")
summary_row(ws, S + 3, "平均脈搏", f"=IF(COUNT({rng_p})=0,\"\",ROUND(AVERAGE({rng_p}),0))", "次/分", "0")
summary_row(
    ws,
    S + 4,
    "判斷",
    f"=IF(B{S}=0,\"\",IF(OR(N(B{S + 1})>=130,N(B{S + 2})>=80),\"平均 ≥ 130/80 mmHg：建議調整生活並諮詢醫師\",\"平均低於 130/80 mmHg\")&IF(B{S}<14,\"（未滿 14 次，僅供參考）\",\"\"))",
    "門檻 130/80 mmHg 來自國健署 722 原則",
)
ws.merge_cells(f"B{S + 4}:N{S + 4}")
widths(ws, [16, 13, 7, 8, 9, 9, 8, 9, 9, 8, 10, 10, 9, 22])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- 臥站血壓
ws = wb.create_sheet("臥站血壓")
title(ws, "臥站血壓（每週 1 次，共 4 次）", "平躺休息 5 分鐘後量，再站起 1 分鐘、3 分鐘各量一次。站起時要有人扶，頭暈就坐下停止。" + SRC_OH, "Q")
labels = ["次數", "日期", "時間", "收縮壓", "舒張壓", "脈搏", "收縮壓", "舒張壓", "脈搏", "收縮壓", "舒張壓", "脈搏", "收縮壓下降", "舒張壓下降", "站起時頭暈", "判斷", "備註"]
header(ws, 4, labels, groups=[(1, 3, "量測"), (4, 6, "平躺 5 分鐘"), (7, 9, "站起 1 分鐘"), (10, 12, "站起 3 分鐘"), (13, 14, "最大下降（自動）"), (15, 15, ""), (16, 16, "自動"), (17, 17, "")])
ex = ["範例", "2026-12-02", "08:00", 132, 80, 68, 118, 76, 80, 120, 78, 78, None, None, "否", None, "範例列不計入"]
for i, v in enumerate(ex, start=1):
    ws.cell(row=5, column=i, value=v)


def oh_formulas(r):
    ws[f"M{r}"] = f"=IF(OR(D{r}=\"\",COUNT(G{r},J{r})=0),\"\",D{r}-MIN(G{r},J{r}))"
    ws[f"N{r}"] = f"=IF(OR(E{r}=\"\",COUNT(H{r},K{r})=0),\"\",E{r}-MIN(H{r},K{r}))"
    ws[f"P{r}"] = f"=IF(M{r}=\"\",\"\",IF(OR(M{r}>=20,N(N{r})>=10),\"符合姿勢性低血壓定義\",\"未達定義\"))"


oh_formulas(5)
style_range(ws, "A5:Q5", font=F_EXAMPLE, fill=FILL_EXAMPLE, align=CENTER)
FIRST, LAST = 6, 9
for k in range(4):
    r = FIRST + k
    ws.cell(row=r, column=1, value=f"第 {k + 1} 週")
    oh_formulas(r)
style_range(ws, f"A{FIRST}:A{LAST}", font=F_BASE, align=CENTER)
style_range(ws, f"B{FIRST}:L{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"O{FIRST}:O{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"Q{FIRST}:Q{LAST}", font=F_INPUT, fill=FILL_INPUT, align=LEFT)
style_range(ws, f"M{FIRST}:P{LAST}", font=F_BASE, align=CENTER)
ws.conditional_formatting.add(f"P{FIRST}:P{LAST}", FormulaRule(formula=[f'P{FIRST}="符合姿勢性低血壓定義"'], fill=FILL_ALERT))
dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"O{FIRST}:O{LAST}")
S = LAST + 2
summary_row(ws, S, "完成次數", f"=COUNT(M{FIRST}:M{LAST})", "共 4 次")
summary_row(ws, S + 1, "收縮壓最大下降", f"=IF(COUNT(M{FIRST}:M{LAST})=0,\"\",MAX(M{FIRST}:M{LAST}))", "mmHg（門檻 ≥ 20）")
summary_row(ws, S + 2, "舒張壓最大下降", f"=IF(COUNT(N{FIRST}:N{LAST})=0,\"\",MAX(N{FIRST}:N{LAST}))", "mmHg（門檻 ≥ 10）")
summary_row(ws, S + 3, "符合定義的次數", f"=COUNTIF(P{FIRST}:P{LAST},\"符合姿勢性低血壓定義\")")
summary_row(ws, S + 4, "站起時頭暈的次數", f"=COUNTIF(O{FIRST}:O{LAST},\"是\")")
widths(ws, [16, 13, 8, 9, 9, 8, 9, 9, 8, 9, 9, 8, 11, 11, 11, 22, 20])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- 夜間如廁
ws = wb.create_sheet("夜間如廁")
title(ws, "夜間如廁（連續 7 晚）", "記錄上床、起床時間與半夜起來上廁所的次數。" + SRC_NOC, "G")
labels = ["天", "日期", "上床時間", "起床時間", "夜間如廁次數", "起身時頭暈", "備註"]
header(ws, 4, labels)
ex = ["範例", "2026-12-01", "22:30", "06:10", 2, "否", "範例列不計入"]
for i, v in enumerate(ex, start=1):
    ws.cell(row=5, column=i, value=v)
style_range(ws, "A5:G5", font=F_EXAMPLE, fill=FILL_EXAMPLE, align=CENTER)
FIRST, LAST = 6, 12
for k in range(7):
    ws.cell(row=FIRST + k, column=1, value=f"第 {k + 1} 晚")
style_range(ws, f"A{FIRST}:A{LAST}", font=F_BASE, align=CENTER)
style_range(ws, f"B{FIRST}:F{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"G{FIRST}:G{LAST}", font=F_INPUT, fill=FILL_INPUT, align=LEFT)
ws.conditional_formatting.add(f"E{FIRST}:E{LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["3"], fill=FILL_ALERT))
dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F{FIRST}:F{LAST}")
S = LAST + 2
rng = f"E{FIRST}:E{LAST}"
summary_row(ws, S, "記錄晚數", f"=COUNT({rng})", "共 7 晚")
summary_row(ws, S + 1, "平均每晚次數", f"=IF(COUNT({rng})=0,\"\",ROUND(AVERAGE({rng}),1))", "次/晚", "0.0")
summary_row(ws, S + 2, "≥ 3 次的晚數", f"=COUNTIF({rng},\">=3\")")
summary_row(ws, S + 3, "起身頭暈的晚數", f"=COUNTIF(F{FIRST}:F{LAST},\"是\")")
summary_row(
    ws,
    S + 4,
    "判斷",
    f"=IF(B{S}=0,\"\",IF(AND(B{S}=7,B{S + 2}=7),\"連續一週每晚 ≥ 3 次：依本設計會提醒就醫\",\"未達提醒條件（連續一週每晚 ≥ 3 次）\"))",
)
widths(ws, [16, 13, 11, 11, 14, 12, 30])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- 洗澡濕度
ws = wb.create_sheet("洗澡濕度")
title(ws, "洗澡時的浴室濕度（7 次）", "用電子溫濕度計記錄。結果用來訂「濕度上升多少就判斷正在洗澡」的門檻。濕度單位 %RH。", "L")
labels = ["次數", "日期", "開始時間", "結束時間", "洗澡分鐘", "洗澡前濕度", "洗澡中最高濕度", "濕度上升", "洗完 10 分鐘後濕度", "開窗或開排風扇", "室溫 °C", "備註"]
header(ws, 4, labels)
ex = ["範例", "2026-12-01", "22:40", "22:58", None, 68, 92, None, 80, "否", 18, "範例列不計入"]
for i, v in enumerate(ex, start=1):
    ws.cell(row=5, column=i, value=v)


def shower_formulas(r):
    # times are typed as hh:mm; MOD handles a shower that crosses midnight
    ws[f"E{r}"] = f"=IF(OR(C{r}=\"\",D{r}=\"\"),\"\",ROUND(MOD(D{r}-C{r},1)*1440,0))"
    ws[f"H{r}"] = f"=IF(COUNT(F{r},G{r})<2,\"\",G{r}-F{r})"


ws["C5"].value = datetime.time(22, 40)
ws["D5"].value = datetime.time(22, 58)
shower_formulas(5)
style_range(ws, "A5:L5", font=F_EXAMPLE, fill=FILL_EXAMPLE, align=CENTER)
ws["C5"].number_format = ws["D5"].number_format = "hh:mm"
FIRST, LAST = 6, 12
for k in range(7):
    r = FIRST + k
    ws.cell(row=r, column=1, value=f"第 {k + 1} 次")
    shower_formulas(r)
style_range(ws, f"A{FIRST}:A{LAST}", font=F_BASE, align=CENTER)
style_range(ws, f"B{FIRST}:D{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"C{FIRST}:D{LAST}", fmt="hh:mm")
style_range(ws, f"F{FIRST}:G{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"I{FIRST}:K{LAST}", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
style_range(ws, f"L{FIRST}:L{LAST}", font=F_INPUT, fill=FILL_INPUT, align=LEFT)
style_range(ws, f"E{FIRST}:E{LAST}", font=F_BASE, align=CENTER)
style_range(ws, f"H{FIRST}:H{LAST}", font=F_BASE, align=CENTER)
dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"J{FIRST}:J{LAST}")
S = LAST + 2
summary_row(ws, S, "記錄次數", f"=COUNT(H{FIRST}:H{LAST})", "共 7 次")
summary_row(ws, S + 1, "平均洗澡分鐘", f"=IF(COUNT(E{FIRST}:E{LAST})=0,\"\",ROUND(AVERAGE(E{FIRST}:E{LAST}),0))", "分鐘")
summary_row(ws, S + 2, "平均濕度上升", f"=IF(COUNT(H{FIRST}:H{LAST})=0,\"\",ROUND(AVERAGE(H{FIRST}:H{LAST}),1))", "%RH", "0.0")
summary_row(ws, S + 3, "最小濕度上升", f"=IF(COUNT(H{FIRST}:H{LAST})=0,\"\",MIN(H{FIRST}:H{LAST}))", "%RH")
summary_row(
    ws,
    S + 4,
    "建議門檻",
    f"=IF(B{S + 3}=\"\",\"\",\"濕度上升 ≥ \"&ROUND(B{S + 3}*0.8,0)&\" %RH 判斷為洗澡\")",
    "假設：取 7 次中最小上升量的 8 成，留一點誤差空間（本設計自訂，期末依資料調整）",
)
widths(ws, [16, 13, 10, 10, 10, 11, 14, 10, 16, 14, 9, 24])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- 摘要
ws = wb.create_sheet("摘要")
ws["A1"] = "實測結果摘要（自動帶入，期末報告第 6 項使用）"
ws["A1"].font = F_TITLE
header(ws, 3, ["項目", "結果", "判斷", "文獻或規則"])
items = [
    ("血壓 7 天平均", "=IF('血壓722'!B22=\"\",\"\",'血壓722'!B22&\" / \"&'血壓722'!B23&\" mmHg\")", "='血壓722'!B25", "國健署 722 原則：居家平均 ≥ 130/80 mmHg 要注意"),
    ("血壓量測完成", "='血壓722'!B21&\" / 14 次\"", "", ""),
    ("臥站血壓最大下降", "=IF('臥站血壓'!B12=\"\",\"\",\"收縮壓 \"&'臥站血壓'!B12&\"、舒張壓 \"&'臥站血壓'!B13&\" mmHg\")", "=IF('臥站血壓'!B11=0,\"\",'臥站血壓'!B14&\" 次符合定義\")", "站起 3 分鐘內下降 ≥ 20／10 mmHg"),
    ("夜間如廁", "=IF('夜間如廁'!B15=\"\",\"\",\"平均 \"&'夜間如廁'!B15&\" 次/晚\")", "='夜間如廁'!B18", "每晚 ≥ 3 次，3 年內跌倒風險較高（Vaughan 2010）"),
    ("洗澡時濕度上升", "=IF('洗澡濕度'!B16=\"\",\"\",\"平均 \"&'洗澡濕度'!B16&\" %RH、最小 \"&'洗澡濕度'!B17&\" %RH\")", "='洗澡濕度'!B18", "本設計的洗澡判斷門檻"),
]
for i, (k, a, b, c) in enumerate(items, start=4):
    ws.cell(row=i, column=1, value=k).font = F_BOLD
    ws.cell(row=i, column=2, value=a).font = F_LINK
    ws.cell(row=i, column=3, value=b).font = F_LINK
    ws.cell(row=i, column=4, value=c).font = F_NOTE
    for col in range(1, 5):
        ws.cell(row=i, column=col).border = BORDER
        ws.cell(row=i, column=col).alignment = LEFT
    ws.row_dimensions[i].height = 36
widths(ws, [20, 38, 46, 46])

# typed dates and times display the same way on every computer
FORMATS = {
    "血壓722": (19, {"B": "yyyy-mm-dd", "D": "hh:mm"}),
    "臥站血壓": (9, {"B": "yyyy-mm-dd", "C": "hh:mm"}),
    "夜間如廁": (12, {"B": "yyyy-mm-dd", "C": "hh:mm", "D": "hh:mm"}),
    "洗澡濕度": (12, {"B": "yyyy-mm-dd", "C": "hh:mm", "D": "hh:mm"}),
}
for name, (last, cols) in FORMATS.items():
    for r in range(5, last + 1):
        wb[name].row_dimensions[r].height = 24  # room to write by hand
    for col, fmt in cols.items():
        for r in range(6, last + 1):
            wb[name][f"{col}{r}"].number_format = fmt

# A4, one page wide, so each sheet can also be printed and filled in by hand
for ws in wb.worksheets:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait" if ws.title == "說明" else "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT)

if "--pdf" in sys.argv:
    # LibreOffice computes the formulas while loading, so the copy needs no recalc first
    del wb["摘要"]
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "measurement-sheets.xlsx")
        wb.save(src)
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--norestore",
             "--convert-to", "pdf", "--outdir", tmp, src],
            check=True, capture_output=True, timeout=180,
        )
        pdf = os.path.join(ROOT, "final", "measurement-sheets.pdf")
        shutil.move(os.path.join(tmp, "measurement-sheets.pdf"), pdf)
    print(pdf)
